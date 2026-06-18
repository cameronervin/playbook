#!/usr/bin/env python3
"""End-to-end smoke test for the local KB service.

The script intentionally reads secrets from the runtime environment/.env only.
It never embeds credentials, emits bearer tokens, or prints full source URLs.
"""
from __future__ import annotations

import argparse
import os
import sys
import tempfile
import time
import uuid
from pathlib import Path
from typing import Any
from urllib.parse import quote

import httpx
from botocore.exceptions import BotoCoreError, ClientError
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


DEFAULT_BASE_URL = "http://localhost:8001"
DEFAULT_QUERY = "Do athletes need to disclose NIL deals?"
DEFAULT_TIMEOUT_SECONDS = 180.0
HTTP_TIMEOUT_SECONDS = 15.0
POLL_INTERVAL_SECONDS = 2.0
SMOKE_TITLE = "KB Smoke NIL Disclosure Test"
SMOKE_FILENAME = "kb-smoke-nil-disclosure.docx"
SMOKE_TEXT = (
    "KB smoke test policy. Athletes must disclose NIL deals before signing. "
    "Compliance review is required before agreements are finalized."
)
PRIVATE_SMOKE_TITLE = "KB Smoke Private Contract Test"
PRIVATE_SMOKE_FILENAME = "kb-smoke-private-contract.docx"
PRIVATE_SMOKE_TEXT = (
    "Private conversation file clause. The uploaded NIL contract requires "
    "department approval before signing."
)
MATCH_TERMS = ("nil", "disclose", "compliance")
PRIVATE_MATCH_TERMS = ("private", "approval", "contract")
TERMINAL_SUCCESS_STATUSES = {"success", "completed", "complete"}
TERMINAL_FAILURE_STATUSES = {"failed", "error"}


class SmokeTestError(RuntimeError):
    """Raised when the smoke test cannot complete successfully."""


def _settings() -> Any:
    from app.core.config import settings  # noqa: PLC0415

    return settings


def _s3_client() -> Any:
    from app.infrastructure.io.s3_client import build_s3_client  # noqa: PLC0415

    return build_s3_client()


def _emit(message: str) -> None:
    sys.stdout.write(f"{message}\n")
    sys.stdout.flush()


def _auth_headers(api_secret: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {api_secret}"}


def _quote_key(key: str) -> str:
    return quote(key, safe="/")


def _build_source_uri(runtime_settings: Any, key: str) -> str:
    bucket = runtime_settings.S3_BUCKET_NAME
    encoded_key = _quote_key(key)
    endpoint = runtime_settings.S3_ENDPOINT_URL.rstrip("/")
    if endpoint:
        return f"{endpoint}/{bucket}/{encoded_key}"
    region = runtime_settings.S3_REGION
    return f"https://{bucket}.s3.{region}.amazonaws.com/{encoded_key}"


def _stage_summary(status_payload: dict[str, Any]) -> str:
    stages = status_payload.get("stages") or []
    parts = []
    for stage in stages:
        stage_name = stage.get("stage") or "unknown"
        stage_status = stage.get("status") or "pending"
        parts.append(f"{stage_name}={stage_status}")
    return ", ".join(parts) if parts else "no stages reported"


def _matching_results(
    search_payload: dict[str, Any],
    *,
    terms: tuple[str, ...] = MATCH_TERMS,
) -> list[dict[str, Any]]:
    matches = []
    for result in search_payload.get("results") or []:
        text = str(result.get("text") or "")
        normalized = text.lower()
        if any(term in normalized for term in terms):
            matches.append(result)
    return matches


def _rerank_smoke_required(runtime_settings: Any) -> bool:
    return (
        getattr(runtime_settings, "KB_SEARCH_STRATEGY", None) == "hybrid"
        and getattr(runtime_settings, "KB_RERANK_ENABLED", False) is True
    )


def _assert_rerank_metadata_if_enabled(
    results: list[dict[str, Any]],
    runtime_settings: Any,
) -> None:
    if not _rerank_smoke_required(runtime_settings):
        return
    if any(
        (result.get("metadata") or {}).get("ranking_strategy") == "hybrid_rerank"
        for result in results
    ):
        return
    raise SmokeTestError(
        "rerank-enabled search returned no result with ranking_strategy=hybrid_rerank"
    )


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run a sanitized end-to-end KB service smoke test."
    )
    parser.add_argument(
        "--base-url",
        default=os.getenv("KB_SMOKE_BASE_URL", DEFAULT_BASE_URL),
        help="KB API base URL. Defaults to KB_SMOKE_BASE_URL or localhost:8001.",
    )
    parser.add_argument(
        "--timeout-seconds",
        type=float,
        default=float(os.getenv("KB_SMOKE_TIMEOUT_SECONDS", DEFAULT_TIMEOUT_SECONDS)),
        help="Maximum seconds to wait for ingestion to finish.",
    )
    parser.add_argument(
        "--query",
        default=os.getenv("KB_SMOKE_QUERY", DEFAULT_QUERY),
        help="Retrieval query to run after ingestion succeeds.",
    )
    parser.add_argument(
        "--keep",
        action="store_true",
        help="Keep the uploaded object and KB document for debugging.",
    )
    parser.add_argument(
        "--include-conversation-file",
        action="store_true",
        help=(
            "Also ingest and search a conversation_file source, verifying it is "
            "excluded from default shared search and included by private scope."
        ),
    )
    return parser.parse_args(argv)


def _create_docx(
    path: Path,
    *,
    title: str = SMOKE_TITLE,
    text: str = SMOKE_TEXT,
) -> None:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(text)
    document.save(path)


def _request_json(
    client: httpx.Client,
    method: str,
    path: str,
    *,
    expected_statuses: set[int],
    headers: dict[str, str] | None = None,
    json: dict[str, Any] | None = None,
) -> dict[str, Any]:
    try:
        response = client.request(method, path, headers=headers, json=json)
        if response.status_code not in expected_statuses:
            detail = _safe_response_detail(response)
            raise SmokeTestError(
                f"{method} {path} returned {response.status_code}: {detail}"
            )
        if response.status_code == httpx.codes.NO_CONTENT:
            return {}
        return response.json()
    except httpx.RequestError as exc:
        raise SmokeTestError(f"{method} {path} request failed: {exc.__class__.__name__}") from exc
    except ValueError as exc:
        raise SmokeTestError(f"{method} {path} returned non-JSON response") from exc


def _safe_response_detail(response: httpx.Response) -> str:
    try:
        payload = response.json()
    except ValueError:
        return response.text[:300]
    detail = payload.get("detail") if isinstance(payload, dict) else payload
    return str(detail)[:300]


def _upload_smoke_docx(path: Path, key: str, runtime_settings: Any) -> int:
    client = _s3_client()
    client.upload_file(str(path), runtime_settings.S3_BUCKET_NAME, key)
    return path.stat().st_size


def _delete_smoke_object(key: str, runtime_settings: Any) -> bool:
    try:
        _s3_client().delete_object(Bucket=runtime_settings.S3_BUCKET_NAME, Key=key)
    except (BotoCoreError, ClientError):
        return False
    else:
        return True


def _delete_document(client: httpx.Client, headers: dict[str, str], document_id: str) -> bool:
    try:
        _request_json(
            client,
            "DELETE",
            f"/api/kb/documents/{document_id}",
            expected_statuses={httpx.codes.NO_CONTENT},
            headers=headers,
        )
    except SmokeTestError:
        return False
    else:
        return True


def _wait_for_ingest_success(
    client: httpx.Client,
    headers: dict[str, str],
    document_id: str,
    timeout_seconds: float,
) -> dict[str, Any]:
    deadline = time.monotonic() + timeout_seconds
    last_summary = ""
    while time.monotonic() < deadline:
        status_payload = _request_json(
            client,
            "GET",
            f"/api/kb/status/documents/{document_id}",
            expected_statuses={httpx.codes.OK},
            headers=headers,
        )
        status = str(status_payload.get("status") or "").lower()
        stage_summary = _stage_summary(status_payload)
        if stage_summary != last_summary:
            _emit(f"  ingest status={status or 'unknown'} stages=({stage_summary})")
            last_summary = stage_summary
        if status in TERMINAL_SUCCESS_STATUSES:
            return status_payload
        if status in TERMINAL_FAILURE_STATUSES:
            error_message = status_payload.get("error_message") or "unknown worker error"
            raise SmokeTestError(
                f"ingestion failed for document {document_id}: {str(error_message)[:300]}"
            )
        time.sleep(POLL_INTERVAL_SECONDS)
    raise SmokeTestError(
        f"timed out after {timeout_seconds:.0f}s waiting for document {document_id}"
    )


def _run_smoke(args: argparse.Namespace) -> None:
    runtime_settings = _settings()
    if not runtime_settings.KB_API_SECRET:
        raise SmokeTestError("KB_API_SECRET is required in the runtime environment")
    if not runtime_settings.S3_BUCKET_NAME:
        raise SmokeTestError("S3_BUCKET_NAME is required in the runtime environment")

    headers = _auth_headers(runtime_settings.KB_API_SECRET)
    organization_id = uuid.uuid4()
    playbook_document_id = uuid.uuid4()
    smoke_run_id = uuid.uuid4()
    s3_key = f"kb-smoke/{smoke_run_id}/{SMOKE_FILENAME}"
    private_s3_key = f"kb-smoke/{smoke_run_id}/{PRIVATE_SMOKE_FILENAME}"
    kb_document_id: str | None = None
    private_kb_document_id: str | None = None
    private_conversation_id = uuid.uuid4()
    private_conversation_file_id = uuid.uuid4()
    cleanup_done = False
    source_uri = _build_source_uri(runtime_settings, s3_key)
    private_source_uri = _build_source_uri(runtime_settings, private_s3_key)

    _emit("KB smoke test starting")
    _emit(f"  api={args.base_url.rstrip('/')}")
    _emit(f"  bucket={runtime_settings.S3_BUCKET_NAME}")
    _emit(f"  object_key={s3_key}")
    if args.include_conversation_file:
        _emit(f"  private_object_key={private_s3_key}")

    with tempfile.TemporaryDirectory(prefix="kb-smoke-") as tmp_dir:
        docx_path = Path(tmp_dir) / SMOKE_FILENAME
        private_docx_path = Path(tmp_dir) / PRIVATE_SMOKE_FILENAME
        _create_docx(docx_path)
        _create_docx(
            private_docx_path,
            title=PRIVATE_SMOKE_TITLE,
            text=PRIVATE_SMOKE_TEXT,
        )

        with httpx.Client(
            base_url=args.base_url.rstrip("/"),
            timeout=HTTP_TIMEOUT_SECONDS,
        ) as client:
            try:
                health = _request_json(
                    client,
                    "GET",
                    "/health",
                    expected_statuses={httpx.codes.OK},
                )
                if health.get("status") != "ok":
                    raise SmokeTestError(
                        f"health check is {health.get('status')}: {health.get('checks')}"
                    )
                _emit("  health=ok")

                config = _request_json(
                    client,
                    "POST",
                    "/api/kb/configuration/resolve",
                    expected_statuses={httpx.codes.OK},
                    headers=headers,
                    json={},
                )
                config_id = config["id"]
                _emit(f"  configuration_id={config_id}")

                size_bytes = _upload_smoke_docx(docx_path, s3_key, runtime_settings)
                _emit(f"  uploaded_bytes={size_bytes}")

                ingest = _request_json(
                    client,
                    "POST",
                    "/api/kb/ingest/document",
                    expected_statuses={httpx.codes.OK, httpx.codes.ACCEPTED},
                    headers=headers,
                    json={
                        "source_type": "admin_upload",
                        "organization_id": str(organization_id),
                        "playbook_document_id": str(playbook_document_id),
                        "configuration_id": config_id,
                        "source_uri": source_uri,
                        "filename": SMOKE_FILENAME,
                        "content_type": (
                            "application/vnd.openxmlformats-officedocument."
                            "wordprocessingml.document"
                        ),
                        "size_bytes": size_bytes,
                        "source_title": SMOKE_TITLE,
                        "is_official": True,
                        "visibility_policy": {"scope": "all_athletes"},
                        "metadata_tags": {"smoke_test": True},
                    },
                )
                kb_document_id = str(ingest["kb_service_document_id"])
                _emit(f"  kb_service_document_id={kb_document_id}")

                _wait_for_ingest_success(
                    client,
                    headers,
                    kb_document_id,
                    args.timeout_seconds,
                )

                search = _request_json(
                    client,
                    "POST",
                    "/api/kb/search",
                    expected_statuses={httpx.codes.OK},
                    headers=headers,
                    json={
                        "query": args.query,
                        "organization_id": str(organization_id),
                        "visibility_context": {"role": "athlete"},
                        "limit": 5,
                        "score_threshold": 0.0,
                    },
                )
                matches = _matching_results(search)
                if not matches:
                    raise SmokeTestError(
                        f"search returned {search.get('total', 0)} results but no smoke-text match"
                    )
                _assert_rerank_metadata_if_enabled(matches, runtime_settings)
                _emit(f"  search_results={search.get('total', len(search.get('results') or []))}")

                private_search_total = None
                if args.include_conversation_file:
                    private_size_bytes = _upload_smoke_docx(
                        private_docx_path,
                        private_s3_key,
                        runtime_settings,
                    )
                    _emit(f"  private_uploaded_bytes={private_size_bytes}")

                    private_ingest = _request_json(
                        client,
                        "POST",
                        "/api/kb/ingest/document",
                        expected_statuses={httpx.codes.OK, httpx.codes.ACCEPTED},
                        headers=headers,
                        json={
                            "source_type": "conversation_file",
                            "organization_id": str(organization_id),
                            "conversation_id": str(private_conversation_id),
                            "conversation_file_id": str(private_conversation_file_id),
                            "configuration_id": config_id,
                            "source_uri": private_source_uri,
                            "filename": PRIVATE_SMOKE_FILENAME,
                            "content_type": (
                                "application/vnd.openxmlformats-officedocument."
                                "wordprocessingml.document"
                            ),
                            "size_bytes": private_size_bytes,
                            "source_title": PRIVATE_SMOKE_TITLE,
                            "visibility_policy": {"scope": "conversation"},
                            "metadata_tags": {"smoke_test": True},
                        },
                    )
                    private_kb_document_id = str(
                        private_ingest["kb_service_document_id"]
                    )
                    _emit(f"  private_kb_service_document_id={private_kb_document_id}")

                    _wait_for_ingest_success(
                        client,
                        headers,
                        private_kb_document_id,
                        args.timeout_seconds,
                    )

                    shared_scope_private_query = _request_json(
                        client,
                        "POST",
                        "/api/kb/search",
                        expected_statuses={httpx.codes.OK},
                        headers=headers,
                        json={
                            "query": "Does the private contract require approval?",
                            "organization_id": str(organization_id),
                            "visibility_context": {"role": "athlete"},
                            "limit": 5,
                            "score_threshold": 0.0,
                        },
                    )
                    leaked_private_matches = _matching_results(
                        shared_scope_private_query,
                        terms=PRIVATE_MATCH_TERMS,
                    )
                    if leaked_private_matches:
                        raise SmokeTestError(
                            "default shared search returned conversation-file text"
                        )

                    private_search = _request_json(
                        client,
                        "POST",
                        "/api/kb/search",
                        expected_statuses={httpx.codes.OK},
                        headers=headers,
                        json={
                            "query": "Does the private contract require approval?",
                            "organization_id": str(organization_id),
                            "source_types": ["conversation_file"],
                            "conversation_id": str(private_conversation_id),
                            "file_ids": [str(private_conversation_file_id)],
                            "limit": 5,
                            "score_threshold": 0.0,
                        },
                    )
                    private_matches = _matching_results(
                        private_search,
                        terms=PRIVATE_MATCH_TERMS,
                    )
                    if not private_matches:
                        raise SmokeTestError(
                            "private conversation-file search returned no smoke-text match"
                        )
                    private_search_total = private_search.get(
                        "total",
                        len(private_search.get("results") or []),
                    )
                    _emit(f"  private_search_results={private_search_total}")

                cleanup_status = "kept"
                if not args.keep:
                    deleted_doc = _delete_document(client, headers, kb_document_id)
                    deleted_private_doc = (
                        True
                        if private_kb_document_id is None
                        else _delete_document(client, headers, private_kb_document_id)
                    )
                    deleted_object = (
                        True
                        if deleted_doc
                        else _delete_smoke_object(s3_key, runtime_settings)
                    )
                    deleted_private_object = (
                        True
                        if private_kb_document_id is None or deleted_private_doc
                        else _delete_smoke_object(private_s3_key, runtime_settings)
                    )
                    cleanup_done = (
                        (deleted_doc or deleted_object)
                        and (deleted_private_doc or deleted_private_object)
                    )
                    cleanup_status = (
                        "deleted" if cleanup_done else "delete_failed"
                    )

                _emit(
                    "PASS "
                    f"document_id={kb_document_id} "
                    f"result_count={search.get('total', len(search.get('results') or []))} "
                    f"private_result_count={private_search_total} "
                    f"cleanup={cleanup_status}"
                )
            finally:
                if not args.keep and not cleanup_done:
                    if kb_document_id is not None:
                        _delete_document(client, headers, kb_document_id)
                    else:
                        _delete_smoke_object(s3_key, runtime_settings)
                    if private_kb_document_id is not None:
                        _delete_document(client, headers, private_kb_document_id)
                    elif args.include_conversation_file:
                        _delete_smoke_object(private_s3_key, runtime_settings)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)
    try:
        _run_smoke(args)
    except SmokeTestError as exc:
        sys.stderr.write(f"FAIL {exc}\n")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
