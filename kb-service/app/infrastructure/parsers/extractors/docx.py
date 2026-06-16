"""Native DOCX parser (python-docx).

The "low complexity" DOCX path: a fast, dependency-light text extractor that
joins all non-empty paragraphs. ``from docx import Document`` is imported lazily
inside the method so this module compiles without python-docx installed.
"""
from __future__ import annotations

from typing import IO

import structlog

from app.infrastructure.parsers.contracts.base import build_text_only_outcome
from app.infrastructure.parsers.contracts.errors import CorruptFileError
from app.infrastructure.parsers.contracts.models import ParseOutcome

logger = structlog.get_logger(__name__)


def _document_text_and_locator(document) -> tuple[str, dict]:
    """Return joined paragraph text and its paragraph range locator."""
    paragraph_entries = [
        (index, paragraph.text.strip())
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    if not paragraph_entries:
        return "", {}

    return "\n".join(text for _index, text in paragraph_entries), {
        "type": "paragraph_range",
        "paragraph_start": paragraph_entries[0][0],
        "paragraph_end": paragraph_entries[-1][0],
    }


class NativeDOCXParser:
    parser_id = "native_docx"

    def parse(self, file: IO[bytes], filename: str) -> list[str]:
        from docx import Document

        try:
            file.seek(0)
            document = Document(file)
            text, _locator = _document_text_and_locator(document)
            if not text:
                return []
        except Exception as exc:  # pragma: no cover - dependency-specific failures
            logger.exception("docx_parse_failed", filename=filename, error=str(exc))
            raise CorruptFileError(f"Failed to parse DOCX: {filename}") from exc
        else:
            return [text]

    def parse_path(self, path: str, filename: str) -> list[str]:
        with open(path, "rb") as file:
            return self.parse(file, filename)

    def parse_outcome_path(self, path: str, filename: str) -> ParseOutcome:
        from docx import Document

        try:
            document = Document(path)
            text, locator = _document_text_and_locator(document)
        except Exception as exc:  # pragma: no cover - dependency-specific failures
            logger.exception("docx_parse_path_failed", filename=filename, error=str(exc))
            raise CorruptFileError(f"Failed to parse DOCX: {filename}") from exc

        return build_text_only_outcome(
            text_segments=[text] if text else [],
            text_segment_locators=[locator] if locator else [],
            selected_parser=self.parser_id,
            route="extractor",
        )
