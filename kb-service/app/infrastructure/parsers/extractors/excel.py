"""Native Excel parser (openpyxl).

Extracts each worksheet as a markdown-ish block: one ``# {sheet name}`` header
followed by pipe-joined non-empty cells per row. ``load_workbook`` is imported
lazily so this module compiles without openpyxl installed.
"""
from __future__ import annotations

from typing import IO

import structlog

from app.infrastructure.parsers.contracts.base import build_text_only_outcome
from app.infrastructure.parsers.contracts.errors import CorruptFileError
from app.infrastructure.parsers.contracts.models import ParseOutcome

logger = structlog.get_logger(__name__)


def _worksheet_segments(workbook) -> list[tuple[str, dict]]:
    """Return worksheet text segments with sheet and row-range locators."""
    sheets: list[tuple[str, dict]] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        rows: list[tuple[int, str]] = []
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            cells = [
                str(cell).strip()
                for cell in row
                if cell is not None and str(cell).strip()
            ]
            if cells:
                rows.append((row_index, " | ".join(cells)))
        if rows:
            sheets.append(
                (
                    f"# {worksheet.title}\n" + "\n".join(text for _row_index, text in rows),
                    {
                        "type": "sheet",
                        "sheet_index": sheet_index,
                        "sheet_name": worksheet.title,
                        "row_start": rows[0][0],
                        "row_end": rows[-1][0],
                    },
                )
            )
    return sheets


class NativeExcelParser:
    parser_id = "native_excel"

    def parse(self, file: IO[bytes], filename: str) -> list[str]:
        from openpyxl import load_workbook

        try:
            file.seek(0)
            workbook = load_workbook(filename=file, read_only=True, data_only=True)
            sheets = [text for text, _locator in _worksheet_segments(workbook)]
        except Exception as exc:  # pragma: no cover - dependency-specific failures
            logger.exception("excel_parse_failed", filename=filename, error=str(exc))
            raise CorruptFileError(f"Failed to parse Excel file: {filename}") from exc
        else:
            workbook.close()
            return sheets

    def parse_path(self, path: str, filename: str) -> list[str]:
        with open(path, "rb") as file:
            return self.parse(file, filename)

    def parse_outcome_path(self, path: str, filename: str) -> ParseOutcome:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(filename=path, read_only=True, data_only=True)
            segments = _worksheet_segments(workbook)
        except Exception as exc:  # pragma: no cover - dependency-specific failures
            logger.exception("excel_parse_path_failed", filename=filename, error=str(exc))
            raise CorruptFileError(f"Failed to parse Excel file: {filename}") from exc
        else:
            workbook.close()

        return build_text_only_outcome(
            text_segments=[text for text, _locator in segments],
            text_segment_locators=[locator for _text, locator in segments],
            selected_parser=self.parser_id,
            route="extractor",
        )
