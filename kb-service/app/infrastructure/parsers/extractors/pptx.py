"""Native PPTX parser (python-pptx).

The "low complexity" PPTX path: one text block per slide built from every shape
that carries text. ``Presentation`` is imported lazily so this module compiles
without python-pptx installed.
"""
from __future__ import annotations

from typing import IO

import structlog

from app.infrastructure.parsers.contracts.base import build_text_only_outcome
from app.infrastructure.parsers.contracts.errors import CorruptFileError
from app.infrastructure.parsers.contracts.models import ParseOutcome

logger = structlog.get_logger(__name__)


def _slide_segments_from_presentation(presentation) -> list[tuple[str, dict]]:
    """Return one text segment per slide with slide locators."""
    slides: list[tuple[str, dict]] = []
    for slide_index, slide in enumerate(presentation.slides):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text:
                    texts.append(text)
        if texts:
            slide_number = slide_index + 1
            slides.append(
                (
                    f"# Slide {slide_number}\n" + "\n".join(texts),
                    {
                        "type": "slide",
                        "slide_index": slide_index,
                        "slide_number": slide_number,
                    },
                )
            )
    return slides


class NativePPTXParser:
    parser_id = "native_pptx"

    def parse(self, file: IO[bytes], filename: str) -> list[str]:
        from pptx import Presentation

        try:
            file.seek(0)
            presentation = Presentation(file)
            slides = [text for text, _locator in _slide_segments_from_presentation(presentation)]
        except Exception as exc:  # pragma: no cover - dependency-specific failures
            logger.exception("pptx_parse_failed", filename=filename, error=str(exc))
            raise CorruptFileError(f"Failed to parse PPTX: {filename}") from exc
        else:
            return slides

    def parse_path(self, path: str, filename: str) -> list[str]:
        with open(path, "rb") as file:
            return self.parse(file, filename)

    def parse_outcome_path(self, path: str, filename: str) -> ParseOutcome:
        from pptx import Presentation

        try:
            presentation = Presentation(path)
            segments = _slide_segments_from_presentation(presentation)
        except Exception as exc:  # pragma: no cover - dependency-specific failures
            logger.exception("pptx_parse_path_failed", filename=filename, error=str(exc))
            raise CorruptFileError(f"Failed to parse PPTX: {filename}") from exc

        return build_text_only_outcome(
            text_segments=[text for text, _locator in segments],
            text_segment_locators=[locator for _text, locator in segments],
            selected_parser=self.parser_id,
            route="extractor",
        )
