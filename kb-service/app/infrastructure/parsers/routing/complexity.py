"""Deterministic document complexity assessment for parser routing.

Classifies a document as ``low``, ``medium``, or ``high`` complexity using
purely deterministic signals and the ``settings.*THRESHOLD*`` knobs — no model
calls, fully reproducible. The router maps these categories to parsers:

    low    -> native extractor (fast, dependency-light)
    medium -> Docling (structured extraction)
    high   -> OCR provider (scanned / image-dominant)

All heavy libraries (``fitz``, ``docx``, ``pptx``) are imported lazily inside
the probe methods so this module compiles without them installed.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from zipfile import BadZipFile

from app.core.config import settings


PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@dataclass(frozen=True)
class ComplexityDecision:
    category: str
    reason_codes: list[str] = field(default_factory=list)
    quality_signals: dict[str, float | int | bool] = field(default_factory=dict)


@dataclass(frozen=True)
class PdfComplexitySignals:
    page_count: int
    empty_page_ratio: float
    ocr_suspected_ratio: float
    chars_per_page_min: int

    def to_dict(self) -> dict[str, float | int]:
        return {
            "page_count": self.page_count,
            "empty_page_ratio": self.empty_page_ratio,
            "ocr_suspected_ratio": self.ocr_suspected_ratio,
            "chars_per_page_min": self.chars_per_page_min,
        }


class ComplexityAssessor:
    """Classifies documents as low, medium, or high complexity."""

    def assess(self, *, path: str, filename: str, mime_type: str) -> ComplexityDecision:
        if mime_type == PDF_MIME:
            return self._assess_pdf(path=path)
        if mime_type == DOCX_MIME:
            return self._assess_docx(path=path)
        if mime_type == PPTX_MIME:
            return self._assess_pptx(path=path)
        return ComplexityDecision(
            category="low",
            reason_codes=["mime_not_complexity_scored"],
            quality_signals={"filename": filename},
        )

    def _assess_pdf(self, *, path: str) -> ComplexityDecision:
        needs_ocr, reasons, signals = self._classify_pdf(path)
        signals_dict = signals.to_dict()
        very_complex_reasons: list[str] = []
        if signals_dict["empty_page_ratio"] >= settings.PDF_VERY_COMPLEX_EMPTY_PAGE_RATIO_THRESHOLD:
            very_complex_reasons.append("very_low_text_density")
        if (
            signals_dict["ocr_suspected_ratio"]
            >= settings.PDF_VERY_COMPLEX_IMAGE_PAGE_RATIO_THRESHOLD
        ):
            very_complex_reasons.append("high_image_dominance")
        if signals_dict["page_count"] >= settings.PDF_VERY_COMPLEX_PAGE_COUNT_THRESHOLD:
            very_complex_reasons.append("high_page_count")

        if very_complex_reasons:
            return ComplexityDecision(
                category="high",
                reason_codes=very_complex_reasons,
                quality_signals=signals_dict,
            )

        if needs_ocr:
            return ComplexityDecision(
                category="medium",
                reason_codes=reasons or ["pdf_complex_signals"],
                quality_signals=signals_dict,
            )

        return ComplexityDecision(
            category="low",
            reason_codes=["pdf_low_complexity"],
            quality_signals=signals_dict,
        )

    @staticmethod
    def _classify_pdf(path: str) -> tuple[bool, list[str], PdfComplexitySignals]:
        """Return (needs_ocr, reason_codes, signals) using deterministic thresholds."""
        with _open_pdf(path) as document:
            page_count = len(document)
            if page_count == 0:
                signals = PdfComplexitySignals(0, 1.0, 1.0, 0)
                return True, ["empty_document"], signals

            empty_pages = 0
            image_dominant_pages = 0
            chars_min = 10**9

            for page in document:
                page_text = page.get_text() or ""
                char_count = len(page_text.strip())
                chars_min = min(chars_min, char_count)
                if char_count < settings.PARSE_OCR_CHARS_THRESHOLD:
                    empty_pages += 1

                page_area = page.rect.width * page.rect.height or 1
                dominant = False
                for image in page.get_images(full=True):
                    bbox = page.get_image_bbox(image)
                    if not bbox:
                        continue
                    ratio = (bbox.width * bbox.height) / page_area
                    if ratio >= settings.PARSE_OCR_IMAGE_RATIO_THRESHOLD:
                        dominant = True
                        break
                if dominant:
                    image_dominant_pages += 1

            empty_ratio = empty_pages / page_count
            image_ratio = image_dominant_pages / page_count
            signals = PdfComplexitySignals(
                page_count=page_count,
                empty_page_ratio=empty_ratio,
                ocr_suspected_ratio=image_ratio,
                chars_per_page_min=chars_min if chars_min != 10**9 else 0,
            )

        reasons: list[str] = []
        if empty_ratio >= settings.PARSE_OCR_EMPTY_PAGE_RATIO_THRESHOLD:
            reasons.append("low_text_density")
        if image_ratio >= settings.PARSE_OCR_IMAGE_PAGE_RATIO_THRESHOLD:
            reasons.append("image_dominant_pages")
        if chars_min < settings.PARSE_OCR_CHARS_THRESHOLD:
            reasons.append("low_chars_page")

        return len(reasons) > 0, reasons, signals

    def _assess_docx(self, *, path: str) -> ComplexityDecision:
        from docx import Document

        try:
            document = Document(path)
        except (ValueError, BadZipFile):
            return ComplexityDecision(
                category="medium",
                reason_codes=["docx_parse_probe_failed"],
                quality_signals={"docx_probe_failed": True},
            )

        paragraph_count = len(document.paragraphs)
        table_count = len(document.tables)
        inline_shape_count = len(document.inline_shapes)
        max_paragraph_chars = max((len((p.text or "").strip()) for p in document.paragraphs), default=0)
        avg_paragraph_chars = int(
            sum(len((p.text or "").strip()) for p in document.paragraphs) / max(paragraph_count, 1)
        )

        quality_signals = {
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "inline_shape_count": inline_shape_count,
            "max_paragraph_chars": max_paragraph_chars,
            "avg_paragraph_chars": avg_paragraph_chars,
        }
        reason_codes: list[str] = []
        if table_count >= settings.DOCX_COMPLEX_TABLE_COUNT_THRESHOLD:
            reason_codes.append("docx_tables_present")
        if inline_shape_count >= settings.DOCX_COMPLEX_INLINE_SHAPE_THRESHOLD:
            reason_codes.append("docx_diagrams_or_images_present")
        if max_paragraph_chars >= settings.DOCX_COMPLEX_MAX_PARAGRAPH_CHARS_THRESHOLD:
            reason_codes.append("docx_dense_paragraphs")

        category = "medium" if reason_codes else "low"
        if not reason_codes:
            reason_codes = ["docx_low_complexity"]

        return ComplexityDecision(
            category=category,
            reason_codes=reason_codes,
            quality_signals=quality_signals,
        )

    def _assess_pptx(self, *, path: str) -> ComplexityDecision:
        from pptx import Presentation

        try:
            presentation = Presentation(path)
        except Exception:
            return ComplexityDecision(
                category="medium",
                reason_codes=["pptx_parse_probe_failed"],
                quality_signals={"pptx_probe_failed": True},
            )
        slide_count = len(presentation.slides)
        shape_count = 0
        table_count = 0
        picture_count = 0
        chart_count = 0
        text_shape_count = 0

        for slide in presentation.slides:
            for shape in slide.shapes:
                shape_count += 1
                if getattr(shape, "has_table", False):
                    table_count += 1
                if getattr(shape, "has_chart", False):
                    chart_count += 1
                if getattr(shape, "shape_type", None) == 13:
                    picture_count += 1
                text = getattr(shape, "text", "")
                if isinstance(text, str) and text.strip():
                    text_shape_count += 1

        quality_signals = {
            "slide_count": slide_count,
            "shape_count": shape_count,
            "table_count": table_count,
            "picture_count": picture_count,
            "chart_count": chart_count,
            "text_shape_count": text_shape_count,
            "avg_shapes_per_slide": round(shape_count / max(slide_count, 1), 2),
        }

        reason_codes: list[str] = []
        if table_count >= settings.PPTX_COMPLEX_TABLE_COUNT_THRESHOLD:
            reason_codes.append("pptx_tables_present")
        if chart_count >= settings.PPTX_COMPLEX_CHART_COUNT_THRESHOLD:
            reason_codes.append("pptx_charts_present")
        if picture_count >= settings.PPTX_COMPLEX_PICTURE_COUNT_THRESHOLD:
            reason_codes.append("pptx_visual_density")
        if shape_count >= settings.PPTX_COMPLEX_SHAPE_COUNT_THRESHOLD:
            reason_codes.append("pptx_dense_layout")

        category = "medium" if reason_codes else "low"
        if not reason_codes:
            reason_codes = ["pptx_low_complexity"]

        return ComplexityDecision(
            category=category,
            reason_codes=reason_codes,
            quality_signals=quality_signals,
        )


def is_docling_supported_mime(mime_type: str) -> bool:
    return mime_type in {PDF_MIME, DOCX_MIME, PPTX_MIME}


def _open_pdf(path: str):
    import fitz

    return fitz.open(filename=path)
